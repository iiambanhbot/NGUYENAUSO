"""
Unit tests cho export/excel_writer.py.
Kiểm tra hàm papers_to_excel_bytes:
- Đọc lại bytes bằng openpyxl
- Kiểm tra sheet 'Tài liệu'
- Kiểm tra 18 tiêu đề cột, hàng tiêu đề in đậm, cố định hàng đầu
- Kiểm tra hyperlink ở cột DOI và URL
- Kiểm tra wrap_text ở các cột nội dung dài
- Kiểm tra danh sách rỗng vẫn trả về file hợp lệ với 1 hàng tiêu đề
"""

from __future__ import annotations

import io

import openpyxl
import pytest
from export.excel_writer import papers_to_excel_bytes
from models.paper import Paper


def test_papers_to_excel_bytes_with_data() -> None:
    """Kiểm tra xuất file Excel khi có danh sách bài báo: tiêu đề, số dòng, định dạng, link."""
    p1 = Paper(
        title="Nghiên cứu kinh tế số 1",
        authors=["Tác giả A", "Tác giả B"],
        year=2023,
        journal="Tạp chí Kinh tế",
        doi="10.1000/econ.01",
        url="https://example.com/p1",
        citations=15,
        abstract="Nội dung abstract bài 1",
        source="OpenAlex",
        relevance_score=92,
        relevance_reason="Rất sát đề tài",
        summary="Tóm tắt sơ lược bài 1",
        methodology="Mô hình VAR",
        data_source="Dữ liệu vĩ mô",
        key_findings="Kinh tế số thúc đẩy năng suất",
        limitations="Mẫu số liệu chưa phân tách khu vực nông thôn",
        policy_implications="Cần hoàn thiện thể chế số",
        notes="Ghi chú quan trọng",
        analyzed=True,
    )
    p2 = Paper(
        title="Nghiên cứu kinh tế số 2",
        authors=["Tác giả C"],
        year=2024,
        journal="Journal of Finance",
        doi="https://doi.org/10.1000/econ.02",
        url="https://example.com/p2",
        citations=8,
        abstract="Nội dung abstract bài 2",
        source="OpenAlex",
        relevance_score=85,
        summary="Tóm lược bài 2",
        key_findings="Tác động lan tỏa tích cực",
        notes="Xem xét thêm",
    )

    excel_bytes = papers_to_excel_bytes([p1, p2])
    assert isinstance(excel_bytes, bytes)
    assert len(excel_bytes) > 0

    wb = openpyxl.load_workbook(io.BytesIO(excel_bytes))
    assert "Tài liệu" in wb.sheetnames

    ws = wb["Tài liệu"]
    assert ws.freeze_panes == "A2"

    # 1 hàng tiêu đề + 2 hàng dữ liệu = 3 hàng
    assert ws.max_row == 3
    assert ws.max_column == 18

    expected_headers = list(Paper().to_excel_row().keys())
    actual_headers = [cell.value for cell in ws[1]]
    assert actual_headers == expected_headers

    # Kiểm tra hàng tiêu đề in đậm
    for cell in ws[1]:
        assert cell.font.bold is True

    # Tìm chỉ số cột DOI, URL, Tóm tắt, Kết quả chính, Hạn chế, Ghi chú
    doi_idx = actual_headers.index("DOI") + 1
    url_idx = actual_headers.index("URL") + 1
    summary_idx = actual_headers.index("Tóm tắt") + 1
    findings_idx = actual_headers.index("Kết quả chính") + 1
    limitations_idx = actual_headers.index("Hạn chế") + 1
    notes_idx = actual_headers.index("Ghi chú") + 1

    # Kiểm tra hyperlink DOI (bài 1: 10.1000/econ.01 -> link https://doi.org/10.1000/econ.01)
    doi_cell = ws.cell(row=2, column=doi_idx)
    assert doi_cell.hyperlink is not None
    assert doi_cell.hyperlink.target == "https://doi.org/10.1000/econ.01"

    # Kiểm tra hyperlink URL
    url_cell = ws.cell(row=2, column=url_idx)
    assert url_cell.hyperlink is not None
    assert url_cell.hyperlink.target == "https://example.com/p1"

    # Kiểm tra wrap text các cột nội dung dài
    for col_idx in (summary_idx, findings_idx, limitations_idx, notes_idx):
        cell = ws.cell(row=2, column=col_idx)
        assert cell.alignment.wrap_text is True


def test_papers_to_excel_bytes_empty_list() -> None:
    """Kiểm tra khi truyền danh sách rỗng vẫn xuất ra file Excel hợp lệ có 18 cột tiêu đề."""
    excel_bytes = papers_to_excel_bytes([])
    assert isinstance(excel_bytes, bytes)
    assert len(excel_bytes) > 0

    wb = openpyxl.load_workbook(io.BytesIO(excel_bytes))
    assert "Tài liệu" in wb.sheetnames

    ws = wb["Tài liệu"]
    assert ws.max_row == 1
    assert ws.max_column == 18

    expected_headers = list(Paper().to_excel_row().keys())
    actual_headers = [cell.value for cell in ws[1]]
    assert actual_headers == expected_headers
    for cell in ws[1]:
        assert cell.font.bold is True
