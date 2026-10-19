"""
Xuất danh sách các bài báo nghiên cứu đã chọn ra định dạng file Excel (.xlsx).
Đầu vào: Danh sách các đối tượng Paper.
Đầu ra: Chuỗi bytes của file Excel định dạng hoàn chỉnh.
"""

from __future__ import annotations

import io
from typing import Any

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
import pandas as pd

from models.paper import Paper


def papers_to_excel_bytes(papers: list[Paper]) -> bytes:
    """
    Chuyển đổi danh sách đối tượng Paper thành dữ liệu file Excel (.xlsx) dạng bytes.

    Đặc điểm định dạng:
    - Sheet 'Tài liệu' sử dụng Paper.to_excel_row (18 cột chuẩn).
    - Hàng tiêu đề in đậm, đóng băng hàng đầu (freeze panes).
    - Cột DOI và URL được gắn hyperlink bấm được.
    - Tự động xuống dòng (wrap text) ở các cột Tóm tắt, Kết quả chính, Hạn chế, Ghi chú,...
    - Độ rộng các cột được thiết lập hợp lý.
    - Danh sách rỗng vẫn trả về file Excel hợp lệ chỉ có hàng tiêu đề.

    Đầu vào:
        papers: Danh sách các đối tượng Paper.

    Đầu ra:
        Chuỗi bytes của file Excel.
    """
    headers = list(Paper().to_excel_row().keys())

    if papers:
        rows = [p.to_excel_row() for p in papers if hasattr(p, "to_excel_row")]
        df = pd.DataFrame(rows)
    else:
        df = pd.DataFrame(columns=headers)

    output = io.BytesIO()

    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        sheet_name = "Tài liệu"
        df.to_excel(writer, sheet_name=sheet_name, index=False)
        ws = writer.sheets[sheet_name]

        # 1. Cố định hàng đầu tiên (freeze header)
        ws.freeze_panes = "A2"

        # 2. Định dạng hàng tiêu đề: in đậm, căn giữa
        header_font = Font(name="Calibri", size=11, bold=True)
        header_fill = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")
        for cell in ws[1]:
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        col_names = [cell.value for cell in ws[1]]

        # Tìm chỉ số các cột liên kết và wrap text
        doi_col_idx = col_names.index("DOI") + 1 if "DOI" in col_names else None
        url_col_idx = col_names.index("URL") + 1 if "URL" in col_names else None

        wrap_cols = {
            "Tóm tắt",
            "Kết quả chính",
            "Hạn chế",
            "Ghi chú",
            "Abstract",
            "Lý do phù hợp",
            "Phương pháp",
            "Dữ liệu nghiên cứu",
            "Hàm ý chính sách",
        }
        wrap_col_indices = {i + 1 for i, name in enumerate(col_names) if name in wrap_cols}

        # Độ rộng hợp lý cho từng cột
        col_widths: dict[str, int] = {
            "Tiêu đề": 35,
            "Tác giả": 25,
            "Năm xuất bản": 14,
            "Tạp chí / Nơi xuất bản": 25,
            "DOI": 22,
            "URL": 30,
            "Số trích dẫn": 13,
            "Nguồn": 12,
            "Abstract": 40,
            "Điểm phù hợp": 14,
            "Lý do phù hợp": 30,
            "Tóm tắt": 40,
            "Phương pháp": 30,
            "Dữ liệu nghiên cứu": 30,
            "Kết quả chính": 40,
            "Hạn chế": 35,
            "Hàm ý chính sách": 30,
            "Ghi chú": 30,
        }

        for i, name in enumerate(col_names, 1):
            col_letter = get_column_letter(i)
            ws.column_dimensions[col_letter].width = col_widths.get(str(name), 20)

        # 3. Định dạng dữ liệu: wrap text và hyperlink
        link_font = Font(name="Calibri", size=11, color="0000FF", underline="single")

        for row in range(2, ws.max_row + 1):
            # Thiết lập wrap text cho các cột nội dung dài
            for col_idx in wrap_col_indices:
                cell = ws.cell(row=row, column=col_idx)
                cell.alignment = Alignment(vertical="top", wrap_text=True)

            # Gắn hyperlink cho cột DOI
            if doi_col_idx:
                cell = ws.cell(row=row, column=doi_col_idx)
                doi_val = str(cell.value or "").strip()
                if doi_val:
                    doi_url = doi_val if doi_val.startswith("http") else f"https://doi.org/{doi_val}"
                    cell.hyperlink = doi_url
                    cell.font = link_font

            # Gắn hyperlink cho cột URL
            if url_col_idx:
                cell = ws.cell(row=row, column=url_col_idx)
                url_val = str(cell.value or "").strip()
                if url_val and url_val.startswith("http"):
                    cell.hyperlink = url_val
                    cell.font = link_font

    return output.getvalue()
